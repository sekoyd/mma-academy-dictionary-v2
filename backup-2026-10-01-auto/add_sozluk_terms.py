"""Merge MMA Sözlük terms into the Dictionary v2 glossary (site JSON + master xlsx).

Inputs (source/sozluk/):
  entries_A/B/C.json  blog-derived records (summarized from Desktop/MMACADEMY/Sözlük/<term>/*.docx)
  acronyms.json       parsed from Sözlük/AAA Music - BAckground/MMA Acronyms.docx
  notion_1/2/3.json   Notion "MMA / Türkiye / Sözlük" terms, rewritten and enriched (not copied)

Idempotent: rows whose Source Document starts with "MMA Sözlük" are dropped and rebuilt.
Run: uv run --with openpyxl python tools/add_sozluk_terms.py
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import TableColumn

ROOT = Path(__file__).resolve().parents[1]
SITE_JSON = ROOT / "site/data/glossary-source.json"
XLSX = ROOT / "source/MMA_Dictionary_Master_Glossary_v1.xlsx"
SOZLUK = ROOT / "source/sozluk"
SOURCE_PREFIX = "MMA Sözlük"

TAG_RETAIL = "Retail Media"
TAG_CONCEPT = "AI & Marketing Concepts"
TAG_MMA = "MMA Programs & Frameworks"
TAG_DIGITAL = "Digital Marketing Acronyms"
NOTION_PAGE = "https://app.notion.com/p/b1869f1c810043c1b50a413180bc6439"
NOTION_NOTE = "Rewritten and enriched from Notion notes plus general knowledge and web sources; editorial review needed."

BLOG_DOCS = {
    "AI Empathy Illusion": "AI Impathy Illusion.docx",
    "AI Orchestrator": "AI Orchestrator.docx",
    "AI Sycophancy": "AI Sycophancy.docx",
    "AI Wrapper": "AI Wrapper.docx",
    "AI-Native Generation": "AI-Native Consumer.docx",
    "Addictive Intelligence": "Addictive Intelligence.docx",
    "Agentic Learner": "agentic learner.txt",
    "Aggragate attention": "Aggragate Attentiıon.docx",
    "Archivetainment": "Archivetainment.docx",
    "Augmented Intelligence": "Augmented Intelligence - 1.docx",
    "Conversional AI Commerce": "Conversational AI Commerce.docx",
    "Creative Intelligence": "Creative Intelligence.docx",
    "Gibberlink": "Gibberlink.docx",
    "Hypertail": "BLOG.docx",
    "LLM SOV": "LLM SOV.docx",
    "Learning Architect": "Learning Architect.docx",
    "Neural Storytelling": "Neural Storytelling.docx",
    "Perpetual Optimism": "Perpetual Optimism.docx",
    "Predictive Copy Performance": "Predictive Copy Performance - BLOG - ENG_TR.docx",
    "Predictive Loyalty": "Predictive Loyalty.docx",
    "Silver Spenders": "Silver Spenders.docx",
    "Supermarketer": "Supermarketer.docx",
    "Synthetic Influence": "Synthetic Influence.docx",
    "Textpectation": "blog_textpectation.docx",
    "The Algorithm Generation": "The Algorithm Generation.docx",
    "Trust Dividend": "Trust Dividend.docx",
    "Webtoon": "Webtoon.docx",
    "Workflow-First": "Workflow-First AI.docx",
    "Workslop": "Workslop.docx",
    "brand stickiness modeling": "Brand Stickiness Modeling.docx",
    "feedforward": "feedforward.docx",
}

# Acronyms already defined in the retail media set; the existing definition is kept.
RETAIL_DUPLICATES = {
    "CTR (Click-Through Rate)", "SOV (Share of Voice)", "CPC (Cost Per Click)",
    "CPA (Cost Per Acquisition)", "CPM (Cost Per Mille)", "ROAS (Return on Ad Spend)",
    "KPI (Key Performance Indicators)", "CDP (Customer Data Platform)",
}
ACRONYM_NOTES = {
    "CTR (Content Traffic Rate)": "Source list expands CTR as 'Content Traffic Rate', which is non-standard; CTR usually means Click-Through Rate (see that entry).",
    "VTR (View-Through Rate)": "Source definition describes view-through conversion; VTR usually means the share of viewers who watch a video ad to completion. Check before publishing.",
}

REVIEW_NOTE = "Summarized from the MMA Sözlük source blog; editorial review needed."

# Hand-written from the Structural Chaos carousel (sc.pdf, image-only) and placeholders for
# folders that contain only images or video.
MANUAL = [
    {
        "Term": "Structural Chaos",
        "Short Definition": "The state in which an organization's complexity grows faster than its coordination, so activity rises while impact falls.",
        "Source Definition / Long Explanation": "Structural chaos does not look like chaos: political shifts, economic pressure, and fragmented markets can look like growth until the system breaks. Its signature is more activity with less impact, because motion is not progress. The response the source proposes is resilience as strategy: anticipate, adapt, recover, align, and create value.",
        "Usage Context / Examples": "From Dr. Ecmel Ayral's presentation at the MMA MEA CMO Board Meeting, framed within the CMO Resilience Playbook.",
        "Synonyms / Related Terms": "CMO Resilience Playbook, Marketing Resilience",
        "TR Definition": "Bir organizasyonda karmaşıklığın koordinasyondan daha hızlı büyümesi; faaliyet artarken etkinin azalması durumu.",
        "Source Document": f"{SOURCE_PREFIX} / Structural Chaos / sc.pdf",
        "Tags / Taxonomy": TAG_CONCEPT,
        "Notes": "Written from the image-only carousel sc.pdf (no text file in the folder); verify wording against the talk.",
    },
    {
        "Term": "BYOAI (Bring Your Own AI)",
        "Short Definition": "The practice of employees using their own, personally chosen AI tools at work without waiting for employer-provided or approved solutions.",
        "Source Definition / Long Explanation": "BYOAI extends the older 'bring your own device' idea to generative AI. The term was popularized by the Microsoft and LinkedIn 2024 Work Trend Index, which found that 75% of knowledge workers used AI at work and 78% of those AI users brought their own tools. For marketing leaders it signals both demand and risk: teams are already experimenting, but unmanaged tools can expose customer and company data and fragment workflows, so the task is to provide sanctioned tools, training, and clear data rules.",
        "Usage Context / Examples": "A content team drafting campaign copy in personal ChatGPT accounts before the company has a licensed, governed AI tool is a typical BYOAI situation.",
        "Synonyms / Related Terms": "Shadow AI, Shadow IT, BYOD, Workflow-First AI",
        "TR Definition": "Çalışanların işverenin sağladığı ya da onayladığı çözümleri beklemeden kendi seçtikleri yapay zeka araçlarını iş yerinde kullanmasıdır.",
        "Source Document": f"{SOURCE_PREFIX} (Notion) / BYOAI; Microsoft & LinkedIn 2024 Work Trend Index",
        "Source Link": "https://news.microsoft.com/source/2024/05/08/microsoft-and-linkedin-release-the-2024-work-trend-index-on-the-state-of-ai-at-work/",
        "Tags / Taxonomy": TAG_CONCEPT,
        "Language": "EN/TR",
        "Content Status": "Draft (enriched)",
        "Notes": NOTION_NOTE,
    },
    *[
        {
            "Term": term,
            "Source Document": f"{SOURCE_PREFIX} / {folder} (images/video only)",
            "Tags / Taxonomy": TAG_CONCEPT,
            "Notes": note,
        }
        for term, folder, note in [
            ("Zebra Consumer", "Zebra Consumer", "Folder has only images; no source text for a definition."),
            ("VIDS", "VIDS", "Folder has only video files and may be a generic media folder rather than a term. Confirm or remove."),
        ]
    ],
]


# Retail rows that had no definition in the source docx; filled editorially on 2026-09-23.
RETAIL_FILLS = {
    "KPI (Key Performance Indicator)": {
        "Short Definition": "A measurable value that shows how well a campaign, channel, or business is progressing toward a defined objective.",
        "Source Definition / Long Explanation": "In retail media, KPIs fall into two groups. Media KPIs are attached to ad delivery, such as ROAS, CTR, sales uplift, and new-to-brand customers. Retail KPIs are the commerce and sales measures used alongside them, such as sales, share of shelf, and basket size. Agreeing on KPIs before a campaign starts makes results comparable across retail media networks.",
        "Synonyms / Related Terms": "Media KPIs, Retail KPIs, ROAS, CTR, New-to-Brand",
        "TR Equivalent": "Temel performans göstergesi (KPI)",
        "TR Definition": "Bir kampanyanın, kanalın ya da işin belirlenen hedefe ne kadar yaklaştığını gösteren ölçülebilir değerdir.",
        "Language": "EN/TR",
        "Notes": "Parent heading in the source docx with no definition of its own; definition written editorially on 2026-09-23 from its sub-entries (Media KPIs, Retail KPIs).",
    },
    "Digital": {
        "Short Definition": "In retail media, the online and screen-based channels where ads run, such as retailer websites, apps, emails, and in-store digital screens, as opposed to printed or physical media.",
        "Source Definition / Long Explanation": "Digital channels let retailers target shoppers with first-party data, change creative in real time, and measure results against actual sales, which is harder with print or static in-store media. Most retail media spend, including sponsored products, onsite display, and offsite extension, runs on digital channels.",
        "Synonyms / Related Terms": "Performance Media, Onsite Retail Media, Offsite Retail Media, DOOH",
        "TR Equivalent": "Dijital",
        "TR Definition": "Perakende medyasında reklamların yayınlandığı web sitesi, uygulama, e-posta ve mağaza içi dijital ekran gibi çevrimiçi ve ekran tabanlı kanallardır.",
        "Language": "EN/TR",
        "Notes": "Orphan line in the source docx (after Performance Media) with no definition; defined editorially on 2026-09-23 in its retail media sense. Consider merging into Performance Media.",
    },
}


# Fast-moving protocols and standards: dated status lines, reviewed by the scheduled agent.
STATUS_HEADER = "Current Status"
MMA_HEADER = "MMA Connection"
WEEK_HEADER = "Featured Week"
STATUS = {
    "MCP (Model Context Protocol)": "As of September 2026: introduced by Anthropic in November 2024; donated in December 2025 to the Agentic AI Foundation under the Linux Foundation, which now governs it as a vendor-neutral open standard.",
    "ACP (Agentic Commerce Protocol)": "As of September 2026: launched by OpenAI and Stripe in September 2025. OpenAI retired Instant Checkout in ChatGPT in March 2026 and moved checkout back to merchants' own sites; the protocol itself remains in active development.",
    "AP2 (Agent Payments Protocol)": "As of September 2026: announced by Google in September 2025 with more than 60 partners, including Mastercard, PayPal and Coinbase.",
    "UCP (Universal Commerce Protocol)": "As of September 2026: announced by Google in January 2026 at NRF, co-developed with Shopify, Etsy, Wayfair, Target and Walmart.",
    "A2A (Agent2Agent Protocol)": "As of September 2026: launched by Google in April 2025 and donated to the Linux Foundation in June 2025; supported by more than 150 partner organizations.",
    "x402 (agentic payments protocol)": "As of September 2026: launched by Coinbase in May 2025; the x402 Foundation was formed with Cloudflare in September 2025, and a multi-chain V2 followed in December 2025.",
    "AAMP (Agentic Advertising Management Protocols)": "As of September 2026: named by IAB Tech Lab in February 2026; AAMP 3.0 is open for public comment until 22 October 2026.",
    "NLWeb (Natural Language Web)": "As of September 2026: introduced by Microsoft at Build 2025; every NLWeb instance also works as an MCP server.",
    "PAIR (Publisher Advertiser Identity Reconciliation)": "As of September 2026: PAIR 1.0 has been an IAB Tech Lab open standard since January 2025.",
    "TAP (Trusted Agent Protocol)": "As of September 2026: launched by Visa with Cloudflare in October 2025 with 12 launch partners, including Shopify, Stripe, Microsoft and Coinbase; the partner list has since grown past 100.",
    "AdCP (Ad Context Protocol)": "As of September 2026: launched in October 2025 by a group including Optable, PubMatic, Scope3, Swivel, Triton Digital and Yahoo; it is built on Anthropic's Model Context Protocol (MCP).",
    "ACS (Agentic Commerce Suite)": "As of September 2026: launched by Stripe in December 2025 through the Stripe Dashboard, its APIs and e-commerce platforms such as Wix, WooCommerce, BigCommerce and Squarespace.",
    "UID2 (Unified ID 2.0)": "As of September 2026: in active use. Google retired its Privacy Sandbox APIs in October 2025, which keeps first-party identity solutions like UID2 relevant.",
}


def letter_of(term: str) -> str:
    first = next((ch for ch in term.upper() if ch.isalpha()), "#")
    return first if "A" <= first <= "Z" else "#"


def blog_rows(headers):
    rows = []
    for name in "ABC":
        for r in json.loads((SOZLUK / f"entries_{name}.json").read_text(encoding="utf-8")):
            doc = BLOG_DOCS[r["folder"]]
            rows.append({
                "Term": r["term"],
                "Language": "EN/TR",
                "Short Definition": r["short_definition"],
                "Source Definition / Long Explanation": r["long_explanation"],
                "Usage Context / Examples": r["usage_context"],
                "Source Document": f"{SOURCE_PREFIX} / {r['folder']} / {doc}",
                "Synonyms / Related Terms": r["related_terms"],
                "Tags / Taxonomy": TAG_CONCEPT,
                "TR Equivalent": r["tr_equivalent"],
                "TR Definition": r["tr_definition"],
                "Content Status": "Draft (summarized from blog)",
                "Notes": f"{REVIEW_NOTE} {r['notes']}" if r["notes"] else REVIEW_NOTE,
            })
    return rows


def notion_rows():
    rows = []
    for name in "123":
        path = SOZLUK / f"notion_{name}.json"
        if not path.exists():
            continue
        for r in json.loads(path.read_text(encoding="utf-8")):
            sources = [u for u in r.get("sources") or [] if u]
            rows.append({
                "Term": r["term"],
                "Language": "EN/TR",
                "Short Definition": r["short_definition"],
                "Source Definition / Long Explanation": r["long_explanation"],
                "Usage Context / Examples": r["usage_context"],
                "Source Document": f"{SOURCE_PREFIX} (Notion) / {r['notion_title'][:80]}",
                "Source Link": " ".join([NOTION_PAGE, *sources]),
                "Synonyms / Related Terms": r["related_terms"],
                "Tags / Taxonomy": r["category"],
                "TR Equivalent": r["tr_equivalent"],
                "TR Definition": r["tr_definition"],
                "Content Status": "Draft (enriched)",
                "Notes": f"{NOTION_NOTE} {r['notes']}" if r.get("notes") else NOTION_NOTE,
            })
    return rows


def research_rows():
    """New 2024-2026 acronyms: web-researched, written, then fact-checked (acronyms_2026*.json)."""
    rows = []
    for path in sorted(SOZLUK.glob("acronyms_2026*.json")):
        for r in json.loads(path.read_text(encoding="utf-8")):
            rows.append({
                "Term": r["term"],
                "Language": "EN/TR",
                "Short Definition": r["short_definition"],
                "Source Definition / Long Explanation": r["long_explanation"],
                "Usage Context / Examples": r["usage_context"],
                "Source Document": f"{SOURCE_PREFIX} (Research 2026)",
                "Source Link": " ".join(u for u in r.get("sources") or [] if u) or None,
                "Synonyms / Related Terms": r["related_terms"],
                "Tags / Taxonomy": r["category"],
                "TR Equivalent": r["tr_equivalent"],
                "TR Definition": r["tr_definition"],
                "Content Status": "Draft (researched 2026-09-23)",
                "Notes": r.get("notes"),
            })
    return rows


def term_of_week_rows():
    """Weekly "Term of the Week" picks (term_of_week.json), written in the MMA Sözlük 4-block style."""
    path = SOZLUK / "term_of_week.json"
    if not path.exists():
        return []
    rows = []
    for r in json.loads(path.read_text(encoding="utf-8")):
        rows.append({
            "Term": r["term"],
            "Language": "EN/TR",
            "Short Definition": r["short_definition"],
            "Source Definition / Long Explanation": r["long_explanation"],
            "Usage Context / Examples": r["usage_context"],
            "Source Document": f"{SOURCE_PREFIX} (Term of the Week)",
            "Source Link": " ".join(u for u in r.get("sources") or [] if u) or None,
            "Synonyms / Related Terms": r.get("related_terms"),
            "Tags / Taxonomy": r["category"],
            "TR Equivalent": r.get("tr_equivalent"),
            "TR Definition": r["tr_definition"],
            "Content Status": f"Term of the Week {r['week']}",
            "Notes": r.get("notes"),
            MMA_HEADER: r.get("mma_connection"),
            WEEK_HEADER: f"{r['week']} | {r['hook']}",
        })
    return rows


def acronym_rows():
    rows = []
    for r in json.loads((SOZLUK / "acronyms.json").read_text(encoding="utf-8")):
        if r["term"] in RETAIL_DUPLICATES:
            continue
        rows.append({
            "Term": r["term"],
            "Language": "EN/TR" if r["tr_def"] else "EN",
            "Short Definition": r["short"],
            "Source Document": f"{SOURCE_PREFIX} / MMA Acronyms.docx",
            "Tags / Taxonomy": r["tag"],
            "TR Equivalent": r["tr_eq"],
            "TR Definition": r["tr_def"],
            "Content Status": "Source list",
            "Notes": ACRONYM_NOTES.get(r["term"]),
        })
    return rows


def main():
    table = json.loads(SITE_JSON.read_text(encoding="utf-8"))
    headers, existing = table[0], table[1:]
    for extra in (STATUS_HEADER, MMA_HEADER, WEEK_HEADER):
        if extra not in headers:
            headers.append(extra)
            existing = [row + [None] for row in existing]
    col = {h: i for i, h in enumerate(headers)}

    kept = [row for row in existing if not str(row[col["Source Document"]] or "").startswith(SOURCE_PREFIX)]
    for row in kept:
        row[col["Tags / Taxonomy"]] = row[col["Tags / Taxonomy"]] or TAG_RETAIL
        for field, value in RETAIL_FILLS.get(row[col["Term"]], {}).items():
            row[col[field]] = value

    new = blog_rows(headers) + MANUAL + notion_rows() + research_rows() + term_of_week_rows() + acronym_rows()
    seen = {}
    for rec in [*({"Term": r[col["Term"]]} for r in kept), *new]:
        key = rec["Term"].lower()
        seen[key] = seen.get(key, 0) + 1
    dupes = [t for t, n in seen.items() if n > 1]
    if dupes:
        print("WARNING duplicate terms:", dupes)
    new_rows = []
    for rec in new:
        rec["Letter"] = letter_of(rec["Term"])
        rec.setdefault("Language", "EN")
        new_rows.append([rec.get(h) or None for h in headers])

    rows = kept + new_rows
    for row in rows:
        row[col[STATUS_HEADER]] = STATUS.get(row[col["Term"]])
    fills_path = SOZLUK / "tr_fills.json"
    if fills_path.exists():
        queue = {}
        for item in json.loads(fills_path.read_text(encoding="utf-8")):
            queue.setdefault(item["term"], []).append(item["tr_definition"])
        for row in rows:
            pending = queue.get(row[col["Term"]])
            if not row[col["TR Definition"]] and pending:
                row[col["TR Definition"]] = pending.pop(0)
                row[col["Language"]] = "EN/TR"
    SITE_JSON.write_text(json.dumps([headers, *rows], ensure_ascii=False, indent=1), encoding="utf-8")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Glossary"]
    ws.delete_rows(2, ws.max_row)
    for index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=index, value=header)
    for row in rows:
        ws.append(row)
    table = ws.tables["MMADictionaryGlossary"]
    for index in range(len(table.tableColumns), len(headers)):
        table.tableColumns.append(TableColumn(id=index + 1, name=headers[index]))
    table.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}"
    wb.save(XLSX)

    counts = {}
    for row in rows:
        counts[row[col["Tags / Taxonomy"]]] = counts.get(row[col["Tags / Taxonomy"]], 0) + 1
    print(f"total={len(rows)} kept={len(kept)} added={len(new_rows)}", counts)


if __name__ == "__main__":
    main()
